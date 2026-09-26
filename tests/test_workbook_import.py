"""Workbook import: the pilot collection surface into ``CollectedDocument``.

Synthetic workbooks carry ``evidence_tier = synthetic_test``. The private pilot
workbook is read only when it is present on disk; it is gitignored and is not
a fixture this file embeds.
"""

from __future__ import annotations

import io
import json
import logging
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import from_excel

from src.collect.workbook import (
    DOCUMENT_HEADERS,
    SEARCH_LOG_HEADERS,
    import_workbook,
)
from src.core.errors import ConfigError
from src.core.ids import author_hash, author_salt_id, doc_id, raw_text_sha256
from src.core.logging import configure_logging
from src.models.enums import EvidenceTier, SourcePlatform

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SALT = "test-salt-phase2"
OTHER_SALT = "test-salt-phase2-other"
COLLECTED = datetime(2026, 9, 26, 12, 17)
AUTHOR = "pilot_user_alpha"
SENTINEL = "ZZ_AUTHOR_SENTINEL_9f3a"
UNICODE_TEXT = (
    "I\u2019m looking\u00a0for the photo.\n\n"
    "I dont want this \u201cfixed\u201d \u2014 the spelling stays."
)

_NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _row(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {
        "source_platform": "reddit",
        "source_type": "post",
        "source_item_id": "item-1",
        "parent_thread_id": None,
        "source_url": "https://www.reddit.com/r/googlephotos/comments/item-1/",
        "source_name": "r/googlephotos",
        "title": "Cannot find a photo",
        "author_name_raw": None,
        "published_at": None,
        "collected_at": COLLECTED,
        "language": "en",
        "raw_text": "I cannot find a photo I know I took.",
        "collection_query": "can't find photo",
        "collection_method": "manual_copy",
        "evidence_tier": EvidenceTier.synthetic_test.value,
        "rating": None,
        "researcher_notes": None,
    }
    row.update(overrides)
    return row


def _search(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {
        "search_session_id": "SEARCH-001",
        "source_platform": "reddit",
        "collection_query": "can't find photo",
        "searched_at": COLLECTED,
        "results_scanned": 3,
        "documents_kept": 1,
        "notes": None,
    }
    row.update(overrides)
    return row


def _write_workbook(
    path: Path,
    document_rows: list[dict[str, object]] | None,
    *,
    search_rows: list[dict[str, object]] | None = None,
    document_headers: tuple[str, ...] | list[str] | None = None,
    search_headers: tuple[str, ...] | list[str] | None = None,
    sheets: tuple[str, ...] = ("documents", "search_log"),
) -> None:
    workbook = Workbook()
    pending = workbook.active
    assert pending is not None
    first = True
    if "documents" in sheets:
        sheet = pending
        sheet.title = "documents"
        first = False
        headers = list(document_headers or DOCUMENT_HEADERS)
        sheet.append(headers)
        for row in document_rows or []:
            sheet.append([row.get(header) for header in headers])
    if "search_log" in sheets:
        sheet = pending if first else workbook.create_sheet("search_log")
        if first:
            sheet.title = "search_log"
        headers = list(search_headers or SEARCH_LOG_HEADERS)
        sheet.append(headers)
        rows = search_rows if search_rows is not None else [_search()]
        if document_rows is not None and search_rows is None:
            rows = [_search(documents_kept=len(document_rows))]
        for row in rows:
            sheet.append([row.get(header) for header in headers])
    workbook.save(path)


def _import(path: Path, output: Path, *, salt: str = SALT):
    return import_workbook(path, author_salt=salt, output_dir=output)


def _jsonl_documents(path: Path) -> list[dict[str, object]]:
    text = path.read_text(encoding="utf-8")
    if not text.strip():
        return []
    return [json.loads(line) for line in text.splitlines()]


def _assert_name_absent(blob: str, name: str, *, where: str) -> None:
    if name and name in blob:
        pytest.fail(f"raw author name leaked into {where}")


@pytest.mark.synthetic
def test_successful_workbook_import_maps_headers_not_positions(tmp_path: Path) -> None:
    """A reversed column order still lands every value on its header."""
    path = tmp_path / "reversed.xlsx"
    output = tmp_path / "out"
    text = "I cannot find a photo I know I took."
    _write_workbook(
        path,
        [
            _row(
                source_item_id="item-1",
                author_name_raw=AUTHOR,
                raw_text=text,
                researcher_notes="kept for context",
                rating=2,
            )
        ],
        document_headers=tuple(reversed(DOCUMENT_HEADERS)),
    )

    first = _import(path, output)
    second = _import(path, tmp_path / "out-again")

    assert first.report.rows_read == 1
    assert first.report.rows_accepted == 1
    assert first.report.rows_rejected == 0
    assert first.report.structure_errors == ()
    assert first.report.search_log_rows_read == 1
    assert first.report.search_log_rows_valid == 1
    document = first.documents[0]
    assert document.raw_text == text
    assert document.source_platform.value == "reddit"
    assert document.source_item_id == "item-1"
    assert document.evidence_tier is EvidenceTier.synthetic_test
    assert document.collection_method.value == "manual_copy"
    assert document.language_reported == "en"
    assert document.rating == 2
    assert document.metadata == {"researcher_notes": "kept for context"}
    assert document.engagement is None
    assert document.doc_id == doc_id("reddit", source_item_id="item-1")
    assert document.doc_id == second.documents[0].doc_id
    assert document.ingest_batch_id == second.documents[0].ingest_batch_id
    assert document.raw_text_sha256 == raw_text_sha256(text)
    stored = _jsonl_documents(first.documents_path)  # type: ignore[arg-type]
    assert stored[0]["raw_text"] == text
    assert stored[0]["source_item_id"] != "SEARCH-001"
    assert "author_name_raw" not in stored[0]
    report_text = first.report_text_path.read_text(encoding="utf-8")
    assert "rows read: 1" in report_text
    assert "rows accepted: 1" in report_text
    assert "not collected documents" in report_text


@pytest.mark.synthetic
def test_raw_text_preserves_unicode_and_paragraph_breaks(tmp_path: Path) -> None:
    path = tmp_path / "unicode.xlsx"
    output = tmp_path / "out"
    _write_workbook(path, [_row(raw_text=UNICODE_TEXT)])

    result = _import(path, output)
    document = result.documents[0]

    assert document.raw_text == UNICODE_TEXT
    assert "\u00a0" in document.raw_text
    assert "\n\n" in document.raw_text
    assert "\u2019" in document.raw_text
    assert "\u201c" in document.raw_text
    assert "\u2014" in document.raw_text
    assert "dont" in document.raw_text
    round_trip = _jsonl_documents(result.documents_path)[0]["raw_text"]  # type: ignore[index]
    assert round_trip == UNICODE_TEXT


@pytest.mark.synthetic
def test_optional_blanks_stay_null(tmp_path: Path) -> None:
    path = tmp_path / "blanks.xlsx"
    output = tmp_path / "out"
    _write_workbook(
        path,
        [
            _row(
                source_item_id="blank-none",
                title=None,
                author_name_raw=None,
                published_at=None,
                parent_thread_id=None,
                rating=None,
                researcher_notes=None,
            ),
            _row(
                source_item_id="blank-space",
                source_url="https://www.reddit.com/r/googlephotos/comments/blank-space/",
                title=None,
                author_name_raw="   ",
                published_at=None,
                parent_thread_id=None,
                rating=None,
                researcher_notes=None,
            ),
        ],
    )

    documents = _import(path, output).documents

    assert {document.source_item_id for document in documents} == {
        "blank-none",
        "blank-space",
    }
    for document in documents:
        assert document.title is None
        assert document.author_hash is None
        assert document.author_salt_id == author_salt_id(SALT)
        assert document.published_at is None
        assert document.parent_thread_id is None
        assert document.rating is None
        assert document.metadata == {}
        assert document.raw_text == "I cannot find a photo I know I took."


@pytest.mark.synthetic
def test_invalid_required_fields_are_rejected_with_row_numbers(tmp_path: Path) -> None:
    path = tmp_path / "missing.xlsx"
    output = tmp_path / "out"
    _write_workbook(
        path,
        [
            _row(source_item_id="missing-text", raw_text=None),
            _row(source_item_id="missing-url", source_url=None),
            _row(source_item_id="missing-collected", collected_at=None),
            _row(source_item_id=None),
        ],
        search_rows=[_search(documents_kept=0)],
    )

    result = _import(path, output)

    assert result.report.rows_read == 4
    assert result.report.rows_accepted == 0
    assert result.report.rows_rejected == 4
    by_row = {item.excel_row: item for item in result.report.rejections}
    assert any("raw_text is required" in reason for reason in by_row[2].reasons)
    assert any("source_url is required" in reason for reason in by_row[3].reasons)
    assert any("collected_at is required" in reason for reason in by_row[4].reasons)
    assert any("source_item_id is required" in reason for reason in by_row[5].reasons)
    assert result.documents_path is not None
    assert _jsonl_documents(result.documents_path) == []


@pytest.mark.synthetic
@pytest.mark.parametrize(
    ("field_name", "bad_value"),
    [
        ("source_platform", "myspace"),
        ("source_type", "tweet"),
        ("collection_method", "clipboard"),
        ("evidence_tier", "hearsay"),
    ],
)
def test_invalid_enums_are_rejected(
    tmp_path: Path, field_name: str, bad_value: str
) -> None:
    path = tmp_path / "enum.xlsx"
    output = tmp_path / "out"
    _write_workbook(
        path,
        [_row(**{field_name: bad_value})],
        search_rows=[_search(documents_kept=0)],
    )

    result = _import(path, output)

    assert result.report.rows_accepted == 0
    assert result.report.rows_rejected == 1
    reason = " ".join(result.report.rejections[0].reasons)
    assert bad_value in reason
    assert "row 2:" in reason
    assert field_name in reason


@pytest.mark.synthetic
def test_invalid_dates_report_the_spreadsheet_row_and_original_value(
    tmp_path: Path,
) -> None:
    serial = 46291.51180555556
    path = tmp_path / "dates.xlsx"
    output = tmp_path / "out"
    _write_workbook(
        path,
        [
            _row(source_item_id="serial", published_at=serial),
            _row(source_item_id="bogus", published_at="not-a-date"),
            _row(source_item_id="impossible", published_at="2026-13-40"),
            _row(
                source_item_id="offset",
                published_at="2026-09-26T08:00:00+05:30",
            ),
        ],
        search_rows=[_search(documents_kept=2)],
    )

    result = _import(path, output)
    by_id = {document.source_item_id: document for document in result.documents}

    assert set(by_id) == {"serial", "offset"}
    assert by_id["serial"].published_at == from_excel(serial).replace(
        tzinfo=timezone.utc
    )
    assert by_id["offset"].published_at == datetime.fromisoformat(
        "2026-09-26T08:00:00+05:30"
    )
    rejected = {item.excel_row: item for item in result.report.rejections}
    bogus = " ".join(rejected[3].reasons)
    impossible = " ".join(rejected[4].reasons)
    assert "row 3:" in bogus
    assert "not-a-date" in bogus
    assert "row 4:" in impossible
    assert "2026-13-40" in impossible
    assert result.report.rows_accepted == 2


@pytest.mark.synthetic
def test_duplicate_source_item_id_rejects_every_copy(tmp_path: Path) -> None:
    path = tmp_path / "dup.xlsx"
    output = tmp_path / "out"
    _write_workbook(
        path,
        [
            _row(source_item_id="dup-1", raw_text="First copy of the same item."),
            _row(
                source_item_id="dup-1",
                source_url="https://www.reddit.com/r/googlephotos/comments/other/",
                raw_text="Second copy of the same item.",
            ),
            _row(source_item_id="ok-1", raw_text="A different item."),
        ],
        search_rows=[_search(documents_kept=1)],
    )

    result = _import(path, output)

    assert result.report.rows_read == 3
    assert result.report.rows_accepted == 1
    assert result.documents[0].source_item_id == "ok-1"
    assert len(result.report.duplicate_ids) == 1
    assert result.report.duplicate_ids[0].source_item_id == "dup-1"
    assert result.report.duplicate_ids[0].excel_rows == (2, 3)
    rejected_rows = {item.excel_row for item in result.report.rejections}
    assert rejected_rows == {2, 3}
    report_text = result.report_text_path.read_text(encoding="utf-8")
    assert "dup-1" in report_text
    assert "Duplicate source_item_id" in report_text


@pytest.mark.synthetic
def test_repeated_thread_url_warns_and_keeps_both_rows(tmp_path: Path) -> None:
    """The reply may be listed before the post. The shared URL is not a rejection."""
    url = "https://support.google.com/photos/thread/389789402/waterfall?hl=en"
    path = tmp_path / "thread.xlsx"
    output = tmp_path / "out"
    _write_workbook(
        path,
        [
            _row(
                source_platform="google_support",
                source_type="forum_reply",
                source_item_id="389789402_op_reply_01",
                parent_thread_id="389789402",
                source_url=url,
                source_name="Google Photos Help Community",
                title=None,
                raw_text="Looking in places narrows it down.",
            ),
            _row(
                source_platform="google_support",
                source_type="support_thread",
                source_item_id="389789402",
                parent_thread_id=None,
                source_url=url,
                source_name="Google Photos Help Community",
                raw_text="Pretty simple I have a photo of a waterfall.",
            ),
        ],
        search_rows=[_search(source_platform="google_support", documents_kept=2)],
    )

    result = _import(path, output)

    assert result.report.rows_accepted == 2
    assert result.report.rows_rejected == 0
    assert result.report.duplicate_ids == ()
    assert len(result.report.repeated_urls) == 1
    repeated = result.report.repeated_urls[0]
    assert repeated.excel_rows == (2, 3)
    assert repeated.source_item_ids == (
        "389789402_op_reply_01",
        "389789402",
    )
    assert any(warning.code == "repeated_source_url" for warning in result.report.warnings)
    by_id = {document.source_item_id: document for document in result.documents}
    assert by_id["389789402_op_reply_01"].parent_thread_id == "389789402"
    assert by_id["389789402"].parent_thread_id is None
    assert by_id["389789402_op_reply_01"].doc_id != by_id["389789402"].doc_id


@pytest.mark.synthetic
def test_parent_thread_must_match_another_source_item_id(tmp_path: Path) -> None:
    path = tmp_path / "parents.xlsx"
    output = tmp_path / "out"
    _write_workbook(
        path,
        [
            _row(source_item_id="post-1", raw_text="The original post."),
            _row(
                source_item_id="child-1",
                parent_thread_id="post-1",
                source_type="comment",
                source_url="https://www.reddit.com/r/googlephotos/comments/child-1/",
                raw_text="A reply on that post.",
            ),
            _row(
                source_item_id="orphan-1",
                parent_thread_id="missing-parent",
                source_url="https://www.reddit.com/r/googlephotos/comments/orphan-1/",
                raw_text="A reply whose parent was not collected.",
            ),
            _row(
                source_item_id="self-1",
                parent_thread_id="self-1",
                source_url="https://www.reddit.com/r/googlephotos/comments/self-1/",
                raw_text="A row that names itself as its parent.",
            ),
        ],
        search_rows=[_search(documents_kept=2)],
    )

    result = _import(path, output)
    accepted = {document.source_item_id for document in result.documents}

    assert accepted == {"post-1", "child-1"}
    reasons = {
        item.source_item_id: " ".join(item.reasons) for item in result.report.rejections
    }
    assert "does not match a source_item_id" in reasons["orphan-1"]
    assert "equals this row's source_item_id" in reasons["self-1"]


@pytest.mark.synthetic
def test_author_hash_is_deterministic_and_uses_the_stored_platform(
    tmp_path: Path,
) -> None:
    path = tmp_path / "author.xlsx"
    _write_workbook(
        path,
        [
            _row(
                source_platform="google_photos_help",
                source_item_id="one",
                author_name_raw=AUTHOR,
            ),
            _row(
                source_platform="google_photos_help",
                source_item_id="two",
                source_url="https://www.reddit.com/r/googlephotos/comments/two/",
                author_name_raw=AUTHOR,
                raw_text="A second post by the same author.",
            ),
        ],
        search_rows=[
            _search(source_platform="google_photos_help", documents_kept=2)
        ],
    )

    first = _import(path, tmp_path / "a")
    again = _import(path, tmp_path / "b")
    other = _import(path, tmp_path / "c", salt=OTHER_SALT)

    expected = author_hash(SALT, SourcePlatform.google_support.value, AUTHOR)
    assert first.documents[0].author_hash == expected
    assert first.documents[1].author_hash == expected
    assert again.documents[0].author_hash == expected
    assert other.documents[0].author_hash == author_hash(
        OTHER_SALT, SourcePlatform.google_support.value, AUTHOR
    )
    assert other.documents[0].author_hash != expected
    assert AUTHOR not in (first.documents[0].author_hash or "")
    assert first.documents[0].author_salt_id == author_salt_id(SALT)
    assert first.documents[0].source_platform is SourcePlatform.google_support
    assert any(
        warning.code == "workbook_platform_alias" and warning.sheet == "documents"
        for warning in first.report.warnings
    )


@pytest.mark.synthetic
def test_raw_author_names_are_absent_from_outputs_and_logs(tmp_path: Path) -> None:
    path = tmp_path / "private.xlsx"
    output = tmp_path / "out"
    _write_workbook(
        path,
        [
            _row(source_item_id="kept", author_name_raw=SENTINEL),
            _row(
                source_item_id="dropped",
                author_name_raw=SENTINEL,
                raw_text=None,
            ),
        ],
        search_rows=[_search(documents_kept=1, notes=f"saw {SENTINEL} while searching")],
    )
    buffer = io.StringIO()
    configure_logging(logging.DEBUG, stream=buffer)

    result = _import(path, output)

    assert result.report.rows_accepted == 1
    assert result.report.rows_rejected == 1
    jsonl = result.documents_path.read_text(encoding="utf-8")  # type: ignore[union-attr]
    report_json = result.report_json_path.read_text(encoding="utf-8")
    report_text = result.report_text_path.read_text(encoding="utf-8")
    logs = buffer.getvalue()
    for blob, where in (
        (jsonl, "collected documents"),
        (report_json, "import report json"),
        (report_text, "import report text"),
        (logs, "application log"),
    ):
        _assert_name_absent(blob, SENTINEL, where=where)
        _assert_name_absent(blob, SALT, where=where)
    assert "author_name_raw" not in jsonl
    assert _REDACTED not in jsonl
    assert _REDACTED in report_json


_REDACTED = "[redacted-author]"


@pytest.mark.synthetic
@pytest.mark.parametrize("missing", ["documents", "search_log"])
def test_missing_required_sheet_accepts_nothing(tmp_path: Path, missing: str) -> None:
    path = tmp_path / "sheets.xlsx"
    output = tmp_path / "out"
    present = tuple(name for name in ("documents", "search_log") if name != missing)
    _write_workbook(path, [_row()] if "documents" in present else None, sheets=present)

    result = _import(path, output)

    assert result.documents == ()
    assert result.report.rows_accepted == 0
    assert result.documents_path is None
    assert any(missing in error for error in result.report.structure_errors)
    assert "missing required sheet" in result.report_text_path.read_text(encoding="utf-8")


@pytest.mark.synthetic
@pytest.mark.parametrize(
    ("sheet", "replacement"),
    [
        ("documents", {"raw_text": "body"}),
        ("search_log", {"searched_at": "when"}),
    ],
)
def test_changed_or_missing_headers_accept_nothing(
    tmp_path: Path, sheet: str, replacement: dict[str, str]
) -> None:
    path = tmp_path / "headers.xlsx"
    output = tmp_path / "out"
    document_headers = list(DOCUMENT_HEADERS)
    search_headers = list(SEARCH_LOG_HEADERS)
    target = document_headers if sheet == "documents" else search_headers
    for old, new in replacement.items():
        target[target.index(old)] = new
    _write_workbook(
        path,
        [_row()],
        document_headers=document_headers,
        search_headers=search_headers,
    )

    result = _import(path, output)

    assert result.documents == ()
    assert result.documents_path is None
    errors = " ".join(result.report.structure_errors)
    assert sheet in errors
    for old, new in replacement.items():
        assert old in errors
        assert new in errors


@pytest.mark.synthetic
def test_search_log_values_are_validated_and_not_imported(tmp_path: Path) -> None:
    path = tmp_path / "audit.xlsx"
    output = tmp_path / "out"
    _write_workbook(
        path,
        [_row()],
        search_rows=[
            _search(
                search_session_id="SEARCH-001",
                searched_at="not-a-date",
                documents_kept=1,
            )
        ],
    )

    result = _import(path, output)

    assert result.report.rows_accepted == 1
    assert result.report.search_log_rows_read == 1
    assert result.report.search_log_rows_valid == 0
    assert result.report.search_log_issues[0].excel_row == 2
    assert "not-a-date" in " ".join(result.report.search_log_issues[0].reasons)
    assert result.documents[0].source_item_id == "item-1"
    assert all(
        document.source_item_id != "SEARCH-001" for document in result.documents
    )


@pytest.mark.synthetic
def test_search_log_rejects_bad_counts_and_stays_out_of_documents(
    tmp_path: Path,
) -> None:
    """Counts are non-negative, and kept cannot exceed scanned.

    A failed audit row is an issue on the report. It is not a document, and it
    does not reject the document sheet.
    """
    assert SEARCH_LOG_HEADERS == (
        "search_session_id",
        "source_platform",
        "collection_query",
        "searched_at",
        "results_scanned",
        "documents_kept",
        "notes",
    )
    path = tmp_path / "counts.xlsx"
    _write_workbook(
        path,
        [_row()],
        search_rows=[
            _search(search_session_id="OK", results_scanned=1, documents_kept=1),
            _search(search_session_id="ZERO", results_scanned=0, documents_kept=0),
            _search(search_session_id="NEGSCAN", results_scanned=-1, documents_kept=0),
            _search(search_session_id="NEGKEPT", results_scanned=1, documents_kept=-4),
            _search(search_session_id="OVER", results_scanned=2, documents_kept=3),
        ],
    )

    result = _import(path, tmp_path / "out")

    assert result.report.rows_accepted == 1
    assert result.report.rows_rejected == 0
    assert result.report.search_log_rows_read == 5
    assert result.report.search_log_rows_valid == 2
    assert [issue.excel_row for issue in result.report.search_log_issues] == [4, 5, 6]
    by_row = {issue.excel_row: " ".join(issue.reasons) for issue in result.report.search_log_issues}
    assert "results_scanned" in by_row[4] and ">= 0" in by_row[4]
    assert "documents_kept" in by_row[5] and ">= 0" in by_row[5]
    assert "exceeds" in by_row[6]
    assert {document.source_item_id for document in result.documents} == {"item-1"}
    assert {entry.search_session_id for entry in result.report.search_log} == {"OK", "ZERO"}
    for entry in result.report.search_log:
        assert entry.results_scanned >= 0
        assert entry.documents_kept >= 0
        assert entry.documents_kept <= entry.results_scanned
        assert entry.searched_at.tzinfo is not None


@pytest.mark.synthetic
def test_empty_salt_is_rejected_before_import(tmp_path: Path) -> None:
    path = tmp_path / "salt.xlsx"
    _write_workbook(path, [_row(author_name_raw=SENTINEL)])

    with pytest.raises(ConfigError, match="AUTHOR_SALT"):
        _import(path, tmp_path / "out", salt="")


@pytest.mark.synthetic
def test_import_does_not_load_phase3_modules() -> None:
    script = """
import sys
import src.collect.workbook
forbidden = [
    name
    for name in sys.modules
    if name.startswith((
        "src.normalize",
        "src.dedupe",
        "src.store",
        "src.pipeline",
        "src.extract",
        "src.models.document_derived",
        "src.models.duplicate_link",
    ))
]
assert not forbidden, forbidden
print("ok")
"""
    result = subprocess.run(
        [sys.executable, "-c", script],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == "ok"


def _pilot_workbook() -> Path | None:
    for folder in ("private", "Private"):
        candidate = (
            PROJECT_ROOT
            / "data"
            / "manual"
            / folder
            / "Google_Photos_Pilot_Collection_Workbook.xlsx"
        )
        if candidate.is_file():
            return candidate
    return None


def _xml_raw_text(path: Path) -> dict[str, str]:
    """``raw_text`` as stored in the xlsx, keyed by ``source_item_id``.

    Reads the worksheet XML directly, and finds columns by header name, so this
    oracle does not share openpyxl's reader and does not depend on column order.
    """
    def load_xml(payload: bytes) -> ET.Element:
        if payload.startswith(b"\xef\xbb\xbf"):
            payload = payload[3:]
        return ET.fromstring(payload)

    with zipfile.ZipFile(path) as archive:
        workbook = load_xml(archive.read("xl/workbook.xml"))
        rels = load_xml(archive.read("xl/_rels/workbook.xml.rels"))
        rel_ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
        targets = {
            node.get("Id"): node.get("Target")
            for node in rels.findall("r:Relationship", rel_ns)
        }
        sheet_id = None
        for sheet in workbook.findall("x:sheets/x:sheet", _NS):
            if sheet.get("name") == "documents":
                sheet_id = sheet.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                )
        if sheet_id is None or sheet_id not in targets:
            raise AssertionError("documents sheet is not in the workbook XML")
        target = targets[sheet_id]
        assert target is not None
        if target.startswith("/"):
            target = target.lstrip("/")
        elif not target.startswith("xl/"):
            target = "xl/" + target
        root = load_xml(archive.read(target))

    def column_of(ref: str) -> str:
        return "".join(character for character in ref if character.isalpha())

    header_row = root.find("x:sheetData/x:row", _NS)
    assert header_row is not None
    columns: dict[str, str] = {}
    for cell in header_row.findall("x:c", _NS):
        node = cell.find("x:v", _NS)
        ref = cell.get("r") or ""
        if node is not None and node.text:
            columns[node.text] = column_of(ref)
    id_column = columns["source_item_id"]
    text_column = columns["raw_text"]

    texts: dict[str, str] = {}
    for row in root.findall("x:sheetData/x:row", _NS)[1:]:
        cells: dict[str, str] = {}
        for cell in row.findall("x:c", _NS):
            node = cell.find("x:v", _NS)
            ref = cell.get("r") or ""
            if node is not None and node.text is not None:
                cells[column_of(ref)] = node.text
        item_id = cells.get(id_column)
        if item_id and text_column in cells:
            texts[item_id] = cells[text_column]
    return texts


def test_pilot_workbook_accepts_seven_documents_and_warns_on_the_shared_url(
    tmp_path: Path,
) -> None:
    path = _pilot_workbook()
    if path is None:
        pytest.skip("private pilot workbook is not in this checkout")

    sheet = load_workbook(path)["documents"]
    header_at = {cell.value: cell.column for cell in sheet[1]}
    authors: list[str] = []
    author_by_item: dict[str, str | None] = {}
    for excel_row in range(2, sheet.max_row + 1):
        item_id = sheet.cell(row=excel_row, column=header_at["source_item_id"]).value
        if not isinstance(item_id, str) or not item_id.strip():
            continue
        author = sheet.cell(row=excel_row, column=header_at["author_name_raw"]).value
        if isinstance(author, str) and author.strip():
            authors.append(author)
            author_by_item[item_id] = author
        else:
            author_by_item[item_id] = None

    assert authors, "pilot workbook has no author_name_raw values to hash"

    buffer = io.StringIO()
    configure_logging(logging.DEBUG, stream=buffer)
    result = _import(path, tmp_path / "pilot")
    other = _import(path, tmp_path / "pilot-other-salt", salt=OTHER_SALT)

    assert result.report.rows_read == 7
    assert result.report.rows_accepted == 7
    assert result.report.rows_rejected == 0
    assert result.report.duplicate_ids == ()
    assert result.report.search_log_rows_read == 1
    assert result.report.search_log_rows_valid == 1
    assert result.report.search_log_issues == ()
    for entry in result.report.search_log:
        assert entry.results_scanned >= 0
        assert entry.documents_kept >= 0
        assert entry.documents_kept <= entry.results_scanned
        assert entry.searched_at.tzinfo is not None
    assert len(result.report.repeated_urls) == 1
    repeated = result.report.repeated_urls[0]
    assert repeated.excel_rows == (7, 8)
    assert set(repeated.source_item_ids) == {
        "389789402",
        "389789402_op_reply_01",
    }
    assert any(warning.code == "repeated_source_url" for warning in result.report.warnings)

    by_id = {document.source_item_id: document for document in result.documents}
    core = {"273154473", "469956273", "48648439", "101084512", "133630115"}
    waterfall = {"389789402", "389789402_op_reply_01"}
    assert set(by_id) == core | waterfall
    xml_text = _xml_raw_text(path)
    for item_id, document in by_id.items():
        assert document.raw_text == xml_text[item_id]
        assert document.raw_text_sha256 == raw_text_sha256(document.raw_text)
        assert document.source_platform is SourcePlatform.google_support
        assert document.evidence_tier is EvidenceTier.direct_user
        assert document.collection_method.value == "manual_copy"
        assert document.author_salt_id == author_salt_id(SALT)
        assert "author_name_raw" not in document.model_dump(mode="json")

    assert by_id["273154473"].published_at is None
    assert by_id["273154473"].author_hash is None
    assert by_id["469956273"].title is not None
    assert by_id["389789402_op_reply_01"].title is None
    assert by_id["389789402_op_reply_01"].parent_thread_id == "389789402"
    assert by_id["389789402"].parent_thread_id is None
    for item_id, author in author_by_item.items():
        document = by_id[item_id]
        if author is None:
            assert document.author_hash is None
        else:
            assert document.author_hash == author_hash(
                SALT, SourcePlatform.google_support.value, author
            )
            _assert_name_absent(
                document.author_hash or "", author, where=f"author_hash for {item_id}"
            )
    other_by_id = {document.source_item_id: document for document in other.documents}
    reply_author = author_by_item["389789402_op_reply_01"]
    post_author = author_by_item["389789402"]
    if reply_author is not None and reply_author == post_author:
        assert by_id["389789402"].author_hash == by_id["389789402_op_reply_01"].author_hash
    for item_id, author in author_by_item.items():
        if author is None:
            assert other_by_id[item_id].author_hash is None
        else:
            assert other_by_id[item_id].author_hash != by_id[item_id].author_hash
            assert other_by_id[item_id].author_hash == author_hash(
                OTHER_SALT, SourcePlatform.google_support.value, author
            )
    assert by_id["389789402_op_reply_01"].metadata.get("researcher_notes")
    assert all(document.rating is None for document in result.documents)

    jsonl_lines = result.documents_path.read_text(encoding="utf-8").splitlines()  # type: ignore[union-attr]
    without_text: list[str] = []
    for line in jsonl_lines:
        payload = json.loads(line)
        assert payload["raw_text"] == by_id[payload["source_item_id"]].raw_text
        payload.pop("raw_text")
        without_text.append(json.dumps(payload))
    blobs = {
        "collected documents": "\n".join(without_text),
        "import report json": result.report_json_path.read_text(encoding="utf-8"),
        "import report text": result.report_text_path.read_text(encoding="utf-8"),
        "application log": buffer.getvalue(),
    }
    full_documents = result.documents_path.read_text(encoding="utf-8")  # type: ignore[union-attr]
    other_documents = other.documents_path.read_text(encoding="utf-8")  # type: ignore[union-attr]
    for where, blob in blobs.items():
        for author in authors:
            _assert_name_absent(blob, author, where=where)
        _assert_name_absent(blob, SALT, where=where)
        _assert_name_absent(blob, OTHER_SALT, where=where)
        assert "author_name_raw" not in blob or where != "collected documents"
    _assert_name_absent(full_documents, SALT, where="collected documents raw text")
    _assert_name_absent(other_documents, OTHER_SALT, where="other-salt documents")
