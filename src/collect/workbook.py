"""Import the analyst pilot workbook into ``CollectedDocument`` records.

The workbook is the collection surface an analyst actually fills in. It is not
a second schema. Two sheets matter:

* ``documents`` — one public item per row, mapped onto ``CollectedDocument``.
* ``search_log`` — how the analyst searched. Validated and reported, never
  turned into a document.

Columns are read by header name. A reordered sheet is the same sheet; a
renamed or missing header is a different one and the import stops before any
row is accepted.

``raw_text`` is copied character for character. Blank optional cells become
null. Missing authors, publication dates, titles, and parent ids stay null —
nothing is inferred from the text or from a neighbouring row.

``author_name_raw`` is hashed at this boundary and then dropped. The name is
not a field on the document, and it is scrubbed from the import report and
from log lines. ``raw_text`` is the exception: it is the verbatim post, and
redacting inside it is Phase 3's length-preserving pass, not import's job.
"""

from __future__ import annotations

import hashlib
import json
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Final

from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict, ValidationError as PydanticValidationError

from src.core.errors import ConfigError, ValidationError
from src.core.ids import (
    author_hash,
    author_salt_id,
    doc_id,
    raw_text_sha256,
    source_url_key,
)
from src.core.logging import get_logger
from src.models.collected_document import CollectedDocument
from src.models.enums import (
    CollectionMethod,
    EvidenceTier,
    SourcePlatform,
    SourceType,
)

_LOG = get_logger("collect.workbook")

DOCUMENT_SHEET: Final[str] = "documents"
SEARCH_LOG_SHEET: Final[str] = "search_log"

#: Exact header spellings. Matched as a set, never as a column index.
DOCUMENT_HEADERS: Final[tuple[str, ...]] = (
    "source_platform",
    "source_type",
    "source_item_id",
    "parent_thread_id",
    "source_url",
    "source_name",
    "title",
    "author_name_raw",
    "published_at",
    "collected_at",
    "language",
    "raw_text",
    "collection_query",
    "collection_method",
    "evidence_tier",
    "rating",
    "researcher_notes",
)

SEARCH_LOG_HEADERS: Final[tuple[str, ...]] = (
    "search_session_id",
    "source_platform",
    "collection_query",
    "searched_at",
    "results_scanned",
    "documents_kept",
    "notes",
)

#: Workbook vocabulary that is not the contract vocabulary.
#:
#: The instructions sheet tells the analyst to enter ``google_photos_help``.
#: Spec Section 15.1 and ``config/sources.yaml`` call that same source
#: ``google_support``: Google Photos Help has no API and is collected by hand.
#: The alias is this one explicit pair. Any other token is rejected. Each use
#: is a warning, so the stored enum is not a silent rewrite.
WORKBOOK_PLATFORM_ALIASES: Final[dict[str, str]] = {
    "google_photos_help": SourcePlatform.google_support.value,
}

#: Keys that must never appear in a written document. ``author_name_raw`` is
#: hashed at the boundary; the other two names are the ways a later edit would
#: most likely put the plaintext back.
_FORBIDDEN_OUTPUT_KEYS: Final[frozenset[str]] = frozenset(
    {"author_name_raw", "author_name", "username"}
)

_REDACTED_AUTHOR: Final[str] = "[redacted-author]"


class CellProblem(Exception):
    """One cell cannot be read. ``reason`` already names the spreadsheet row."""

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


class ImportWarning(BaseModel):
    """One non-fatal finding. Repeated thread URLs are warnings, not rejections."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    code: str
    message: str
    sheet: str
    excel_rows: tuple[int, ...] = ()
    source_item_ids: tuple[str, ...] = ()
    source_url: str | None = None


class DuplicateSourceItem(BaseModel):
    """A ``source_item_id`` that occurs on more than one document row."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    source_item_id: str
    excel_rows: tuple[int, ...]


class RepeatedSourceUrl(BaseModel):
    """A source URL shared by more than one document row.

    Sharing a URL does not reject the rows. A thread and the original poster's
    reply on that thread legitimately carry the same thread-level permalink
    when a comment-level link was not captured; ``source_item_id`` and
    ``parent_thread_id`` are what distinguish them.
    """

    model_config = ConfigDict(extra="forbid", frozen=True)

    source_url: str
    excel_rows: tuple[int, ...]
    source_item_ids: tuple[str, ...]


class RowRejection(BaseModel):
    """Why one spreadsheet row was not accepted."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    excel_row: int
    sheet: str
    source_item_id: str | None = None
    reasons: tuple[str, ...]


class AcceptedRow(BaseModel):
    """Pointer from a spreadsheet row to the document it became.

    Identifiers only. The text lives in the documents file, and the author
    name lives nowhere in the output.
    """

    model_config = ConfigDict(extra="forbid", frozen=True)

    excel_row: int
    source_item_id: str
    doc_id: str


class SearchLogSummary(BaseModel):
    """One validated search-log row. Not a ``CollectedDocument``."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    excel_row: int
    search_session_id: str
    source_platform: str
    collection_query: str
    searched_at: datetime
    results_scanned: int
    documents_kept: int
    notes: str | None = None


class ImportReport(BaseModel):
    """Machine-readable result of one workbook import."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    workbook_name: str
    workbook_sha256: str
    ingest_batch_id: str
    rows_read: int
    rows_accepted: int
    rows_rejected: int
    warnings: tuple[ImportWarning, ...]
    duplicate_ids: tuple[DuplicateSourceItem, ...]
    repeated_urls: tuple[RepeatedSourceUrl, ...]
    rejections: tuple[RowRejection, ...]
    accepted_rows: tuple[AcceptedRow, ...]
    search_log_rows_read: int
    search_log_rows_valid: int
    search_log_issues: tuple[RowRejection, ...]
    search_log: tuple[SearchLogSummary, ...]
    structure_errors: tuple[str, ...]


@dataclass(frozen=True)
class WorkbookImport:
    """Documents plus the report that was written beside them."""

    documents: tuple[CollectedDocument, ...]
    report: ImportReport
    output_dir: Path
    documents_path: Path | None
    report_json_path: Path
    report_text_path: Path


@dataclass
class _DocumentRow:
    excel_row: int
    source_item_id: str | None = None
    source_url: str | None = None
    parent_thread_id: str | None = None
    platform_alias: str | None = None
    reasons: list[str] = field(default_factory=list)
    document: CollectedDocument | None = None


@dataclass
class _SearchRow:
    excel_row: int
    summary: SearchLogSummary | None = None
    reasons: list[str] = field(default_factory=list)
    platform_alias: str | None = None
    documents_kept: int | None = None


def import_workbook(
    path: Path | str,
    *,
    author_salt: str,
    output_dir: Path | str,
) -> WorkbookImport:
    """Import ``documents`` into ``CollectedDocument`` records and write a report.

    ``author_salt`` is the HMAC key. It is required even when every author cell
    is blank, because ``author_salt_id`` is stored on every document and an
    empty salt must never be allowed to look like a hash. The caller passes
    the salt in. This module does not read the process environment; that stays
    in ``src.core.config``.

    Raises ``ConfigError`` when ``author_salt`` is empty. A missing file
    raises ``FileNotFoundError``. A file that is not a workbook, or a workbook
    whose sheets or headers do not match, still writes a report and accepts
    zero documents.
    """
    if not author_salt:
        raise ConfigError(
            "AUTHOR_SALT is not set but is required for author hashing at "
            "import. Add it to .env (see .env.example)."
        )

    source = Path(path)
    destination = Path(output_dir)
    file_bytes = source.read_bytes()
    digest = hashlib.sha256(file_bytes).hexdigest()
    batch_id = f"manual-{digest[:12]}"
    salt_id = author_salt_id(author_salt)
    # Author display names are added while rows are read. The salt is included
    # from the start so a report or log line cannot echo it. This set is applied
    # to the report and the log only; raw_text is written unchanged.
    names: set[str] = {author_salt}

    try:
        workbook = _open_workbook(file_bytes)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError):
        report = _scrub_report(
            _empty_report(
                source.name,
                digest,
                batch_id,
                structure_errors=("workbook is not a readable .xlsx file",),
            ),
            names,
        )
        written = _write_outputs(destination, report, ())
        _log_report(report, names)
        return WorkbookImport((), report, destination, None, written[0], written[1])

    try:
        report, documents = _import_sheets(
            workbook,
            workbook_name=source.name,
            workbook_sha256=digest,
            batch_id=batch_id,
            author_salt=author_salt,
            salt_id=salt_id,
            author_names=names,
        )
    finally:
        workbook.close()

    report = _scrub_report(report, names)
    written = _write_outputs(destination, report, documents)
    _log_report(report, names)
    return WorkbookImport(
        documents,
        report,
        destination,
        written[2],
        written[0],
        written[1],
    )


def _open_workbook(file_bytes: bytes) -> Workbook:
    return load_workbook_from_bytes(file_bytes)


def load_workbook_from_bytes(file_bytes: bytes) -> Workbook:
    """Open an xlsx from bytes already read for hashing.

    ``data_only=False`` keeps the cell value rather than a cached formula
    result. The pilot sheets store text and date serials directly, and this
    flag is what keeps ``raw_text`` on the value that was typed.
    """
    from openpyxl import load_workbook

    return load_workbook(BytesIO(file_bytes), data_only=False, rich_text=True)


def _import_sheets(
    workbook: Workbook,
    *,
    workbook_name: str,
    workbook_sha256: str,
    batch_id: str,
    author_salt: str,
    salt_id: str,
    author_names: set[str],
) -> tuple[ImportReport, tuple[CollectedDocument, ...]]:
    structure: list[str] = []
    documents_sheet = _require_sheet(workbook, DOCUMENT_SHEET, structure)
    search_sheet = _require_sheet(workbook, SEARCH_LOG_SHEET, structure)
    document_headers = (
        _require_headers(documents_sheet, DOCUMENT_HEADERS, DOCUMENT_SHEET, structure)
        if documents_sheet is not None
        else None
    )
    search_headers = (
        _require_headers(search_sheet, SEARCH_LOG_HEADERS, SEARCH_LOG_SHEET, structure)
        if search_sheet is not None
        else None
    )
    if structure or document_headers is None or search_headers is None:
        return (
            _empty_report(
                workbook_name,
                workbook_sha256,
                batch_id,
                structure_errors=tuple(structure),
            ),
            (),
        )

    assert documents_sheet is not None and search_sheet is not None
    parsed = [
        _parse_document_row(
            excel_row,
            values,
            batch_id=batch_id,
            author_salt=author_salt,
            salt_id=salt_id,
            author_names=author_names,
        )
        for excel_row, values in _data_rows(
            documents_sheet, document_headers, text_fields=frozenset({"raw_text"})
        )
    ]
    search_rows = [
        _parse_search_row(excel_row, values)
        for excel_row, values in _data_rows(search_sheet, search_headers)
    ]

    duplicates = _apply_duplicate_ids(parsed)
    _apply_parent_links(parsed)
    repeated = _repeated_urls(parsed)
    warnings = _warnings(parsed, search_rows, repeated)
    documents, accepted, rejections = _partition_documents(parsed)
    search_issues = tuple(
        RowRejection(
            excel_row=row.excel_row,
            sheet=SEARCH_LOG_SHEET,
            reasons=tuple(row.reasons),
        )
        for row in search_rows
        if row.reasons
    )
    search_valid = tuple(row.summary for row in search_rows if row.summary is not None and not row.reasons)

    if len(documents) + len(rejections) != len(parsed):
        raise ValidationError(
            "import partitioned rows incorrectly: accepted plus rejected "
            f"is {len(documents) + len(rejections)}, rows read is {len(parsed)}"
        )

    report = ImportReport(
        workbook_name=workbook_name,
        workbook_sha256=workbook_sha256,
        ingest_batch_id=batch_id,
        rows_read=len(parsed),
        rows_accepted=len(documents),
        rows_rejected=len(rejections),
        warnings=tuple(warnings),
        duplicate_ids=tuple(duplicates),
        repeated_urls=tuple(repeated),
        rejections=tuple(rejections),
        accepted_rows=tuple(accepted),
        search_log_rows_read=len(search_rows),
        search_log_rows_valid=len(search_valid),
        search_log_issues=search_issues,
        search_log=search_valid,
        structure_errors=(),
    )
    return report, documents


def _require_sheet(
    workbook: Workbook, name: str, structure: list[str]
) -> Worksheet | None:
    if name not in workbook.sheetnames:
        structure.append(f"missing required sheet {name!r}")
        return None
    return workbook[name]


def _require_headers(
    sheet: Worksheet,
    expected: tuple[str, ...],
    sheet_name: str,
    structure: list[str],
) -> dict[str, int] | None:
    found: dict[str, int] = {}
    for cell in sheet[1]:
        value = cell.value
        if value is None:
            continue
        if not isinstance(value, str):
            structure.append(
                f"{sheet_name} sheet header {cell.coordinate} is not text"
            )
            continue
        if value in found:
            structure.append(f"{sheet_name} sheet repeats column {value!r}")
            continue
        found[value] = cell.column

    expected_set = set(expected)
    missing = [name for name in expected if name not in found]
    unexpected = sorted(name for name in found if name not in expected_set)
    if missing:
        structure.append(
            f"{sheet_name} sheet is missing column(s): {', '.join(missing)}"
        )
    if unexpected:
        structure.append(
            f"{sheet_name} sheet has unexpected column(s): {', '.join(unexpected)}"
        )
    if missing or unexpected or any(message.startswith(f"{sheet_name} sheet") for message in structure):
        return None
    return found


def _data_rows(
    sheet: Worksheet,
    headers: dict[str, int],
    *,
    text_fields: frozenset[str] = frozenset(),
) -> list[tuple[int, dict[str, Any]]]:
    """Non-empty rows under the header, keyed by header name.

    The pilot template styles fifty rows in advance. A styled row with no
    values is not a document, and it is not a rejection either. ``raw_text``
    is empty only when the cell is missing or ``""``; a whitespace-only post
    is still a row, so it can be rejected instead of disappearing.
    """
    rows: list[tuple[int, dict[str, Any]]] = []
    for excel_row in range(2, sheet.max_row + 1):
        values = {
            name: sheet.cell(row=excel_row, column=column).value
            for name, column in headers.items()
        }
        if _row_has_data(values, text_fields=text_fields):
            rows.append((excel_row, values))
    return rows


def _row_has_data(values: dict[str, Any], *, text_fields: frozenset[str]) -> bool:
    for name, value in values.items():
        if name in text_fields:
            if isinstance(value, str):
                if value != "":
                    return True
            elif value is not None:
                return True
        elif not _is_blank(value):
            return True
    return False


def _is_blank(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    return False


def _parse_document_row(
    excel_row: int,
    values: dict[str, Any],
    *,
    batch_id: str,
    author_salt: str,
    salt_id: str,
    author_names: set[str],
) -> _DocumentRow:
    parsed = _DocumentRow(excel_row=excel_row)
    reasons = parsed.reasons

    platform, alias = _take_enum(
        values.get("source_platform"),
        SourcePlatform,
        "source_platform",
        excel_row,
        reasons,
        aliases=WORKBOOK_PLATFORM_ALIASES,
    )
    parsed.platform_alias = alias
    source_type, _ = _take_enum(
        values.get("source_type"), SourceType, "source_type", excel_row, reasons
    )
    evidence_tier, _ = _take_enum(
        values.get("evidence_tier"),
        EvidenceTier,
        "evidence_tier",
        excel_row,
        reasons,
    )
    collection_method, _ = _take_enum(
        values.get("collection_method"),
        CollectionMethod,
        "collection_method",
        excel_row,
        reasons,
    )

    parsed.source_item_id = _take_identifier(
        values.get("source_item_id"), "source_item_id", excel_row, reasons, required=True
    )
    parsed.parent_thread_id = _take_identifier(
        values.get("parent_thread_id"),
        "parent_thread_id",
        excel_row,
        reasons,
        required=False,
    )
    parsed.source_url = _take_url(values.get("source_url"), excel_row, reasons)
    source_name = _take_text(
        values.get("source_name"), "source_name", excel_row, reasons, required=True
    )
    title = _take_text(values.get("title"), "title", excel_row, reasons, required=False)
    language = _take_text(
        values.get("language"), "language", excel_row, reasons, required=True
    )
    collection_query = _take_text(
        values.get("collection_query"),
        "collection_query",
        excel_row,
        reasons,
        required=True,
    )
    notes = _take_text(
        values.get("researcher_notes"),
        "researcher_notes",
        excel_row,
        reasons,
        required=False,
    )
    raw_text = _take_raw_text(values.get("raw_text"), excel_row, reasons)
    published_at = _take_datetime(
        values.get("published_at"), "published_at", excel_row, reasons, required=False
    )
    collected_at = _take_datetime(
        values.get("collected_at"), "collected_at", excel_row, reasons, required=True
    )
    rating = _take_rating(values.get("rating"), excel_row, reasons)
    author = _remember_author(values.get("author_name_raw"), excel_row, reasons, author_names)

    if reasons:
        return parsed
    assert platform is not None
    assert source_type is not None
    assert evidence_tier is not None
    assert collection_method is not None
    assert parsed.source_item_id is not None
    assert parsed.source_url is not None
    assert source_name is not None
    assert language is not None
    assert collection_query is not None
    assert raw_text is not None
    assert collected_at is not None

    hashed = author_hash(author_salt, platform.value, author) if author is not None else None
    metadata: dict[str, Any] = {}
    if notes is not None:
        metadata["researcher_notes"] = notes

    try:
        parsed.document = CollectedDocument(
            doc_id=doc_id(platform.value, source_item_id=parsed.source_item_id),
            ingest_batch_id=batch_id,
            source_platform=platform,
            source_type=source_type,
            evidence_tier=evidence_tier,
            source_item_id=parsed.source_item_id,
            parent_thread_id=parsed.parent_thread_id,
            source_url=parsed.source_url,  # type: ignore[arg-type]
            source_url_key=source_url_key(parsed.source_url),
            source_name=source_name,
            title=title,
            author_hash=hashed,
            author_salt_id=salt_id,
            published_at=published_at,
            collected_at=collected_at,
            language_reported=language,
            raw_text=raw_text,
            raw_text_sha256=raw_text_sha256(raw_text),
            collection_query=collection_query,
            collection_method=collection_method,
            rating=rating,
            engagement=None,
            metadata=metadata,
        )
    except PydanticValidationError as exc:
        reasons.extend(_pydantic_reasons(exc, excel_row))
        parsed.document = None
    except ValueError as exc:
        reasons.append(f"row {excel_row}: {exc}")
        parsed.document = None
    return parsed


def _parse_search_row(excel_row: int, values: dict[str, Any]) -> _SearchRow:
    parsed = _SearchRow(excel_row=excel_row)
    reasons = parsed.reasons
    platform, alias = _take_enum(
        values.get("source_platform"),
        SourcePlatform,
        "source_platform",
        excel_row,
        reasons,
        aliases=WORKBOOK_PLATFORM_ALIASES,
    )
    parsed.platform_alias = alias
    session = _take_text(
        values.get("search_session_id"),
        "search_session_id",
        excel_row,
        reasons,
        required=True,
    )
    query = _take_text(
        values.get("collection_query"),
        "collection_query",
        excel_row,
        reasons,
        required=True,
    )
    notes = _take_text(values.get("notes"), "notes", excel_row, reasons, required=False)
    searched_at = _take_datetime(
        values.get("searched_at"), "searched_at", excel_row, reasons, required=True
    )
    scanned = _take_whole_number(
        values.get("results_scanned"), "results_scanned", excel_row, reasons
    )
    kept = _take_whole_number(
        values.get("documents_kept"), "documents_kept", excel_row, reasons
    )
    parsed.documents_kept = kept
    if scanned is not None and kept is not None and kept > scanned:
        reasons.append(
            f"row {excel_row}: documents_kept ({kept}) exceeds "
            f"results_scanned ({scanned})"
        )
    if reasons:
        return parsed
    assert platform is not None
    assert session is not None
    assert query is not None
    assert searched_at is not None
    assert scanned is not None
    assert kept is not None
    parsed.summary = SearchLogSummary(
        excel_row=excel_row,
        search_session_id=session,
        source_platform=platform.value,
        collection_query=query,
        searched_at=searched_at,
        results_scanned=scanned,
        documents_kept=kept,
        notes=notes,
    )
    return parsed


def _take_enum(
    value: object,
    enum_cls: type[SourcePlatform] | type[SourceType] | type[EvidenceTier] | type[CollectionMethod],
    field_name: str,
    excel_row: int,
    reasons: list[str],
    *,
    aliases: dict[str, str] | None = None,
) -> tuple[Any, str | None]:
    try:
        if _is_blank(value):
            raise CellProblem(f"row {excel_row}: {field_name} is required")
        if not isinstance(value, str):
            raise CellProblem(f"row {excel_row}: {field_name} must be text")
        token = value.strip()
        stored = (aliases or {}).get(token, token)
        alias = token if stored != token else None
        try:
            parsed = enum_cls(stored)
        except ValueError:
            allowed = ", ".join(member.value for member in enum_cls)
            raise CellProblem(
                f"row {excel_row}: {field_name} {token!r} is not one of: {allowed}"
            ) from None
        return parsed, alias
    except CellProblem as exc:
        reasons.append(exc.reason)
        return None, None


def _take_identifier(
    value: object,
    field_name: str,
    excel_row: int,
    reasons: list[str],
    *,
    required: bool,
) -> str | None:
    try:
        if _is_blank(value):
            if required:
                raise CellProblem(f"row {excel_row}: {field_name} is required")
            return None
        if isinstance(value, bool):
            raise CellProblem(f"row {excel_row}: {field_name} must be text")
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float):
            if not value.is_integer():
                raise CellProblem(
                    f"row {excel_row}: {field_name} is not a whole identifier ({value!r})"
                )
            return str(int(value))
        if isinstance(value, str):
            return value
        raise CellProblem(f"row {excel_row}: {field_name} must be text")
    except CellProblem as exc:
        reasons.append(exc.reason)
        return None


def _take_text(
    value: object,
    field_name: str,
    excel_row: int,
    reasons: list[str],
    *,
    required: bool,
) -> str | None:
    try:
        if isinstance(value, str):
            if value.strip() == "":
                value = None
            else:
                return value
        if _is_blank(value):
            if required:
                raise CellProblem(f"row {excel_row}: {field_name} is required")
            return None
        raise CellProblem(f"row {excel_row}: {field_name} must be text")
    except CellProblem as exc:
        reasons.append(exc.reason)
        return None


def _take_url(value: object, excel_row: int, reasons: list[str]) -> str | None:
    try:
        if _is_blank(value):
            raise CellProblem(f"row {excel_row}: source_url is required")
        if not isinstance(value, str):
            raise CellProblem(f"row {excel_row}: source_url must be text")
        url = value.strip()
        if not url:
            raise CellProblem(f"row {excel_row}: source_url is required")
        return url
    except CellProblem as exc:
        reasons.append(exc.reason)
        return None


def _take_raw_text(value: object, excel_row: int, reasons: list[str]) -> str | None:
    """Copy the cell exactly. Spelling, punctuation, ``\\xa0``, and breaks stay."""
    try:
        if isinstance(value, CellRichText):
            text = _rich_text(value)
        elif isinstance(value, str):
            text = value
        elif value is None:
            raise CellProblem(f"row {excel_row}: raw_text is required")
        else:
            raise CellProblem(f"row {excel_row}: raw_text must be text")
        if text == "":
            raise CellProblem(f"row {excel_row}: raw_text is required")
        if not text.strip():
            raise CellProblem(
                f"row {excel_row}: raw_text must contain non-whitespace characters"
            )
        return text
    except CellProblem as exc:
        reasons.append(exc.reason)
        return None


def _rich_text(value: CellRichText) -> str:
    parts: list[str] = []
    for block in value:
        if isinstance(block, str):
            parts.append(block)
        else:
            parts.append(block.text)
    return "".join(parts)


def _take_datetime(
    value: object,
    field_name: str,
    excel_row: int,
    reasons: list[str],
    *,
    required: bool,
) -> datetime | None:
    """Parse a date cell. Blank optional dates stay null; bad dates do not.

    Turning an unreadable date into null would claim the analyst left it
    blank. The row is rejected, and the reason quotes the original cell and
    the spreadsheet row.

    Excel serials and naive datetimes have no timezone. The contract requires
    an aware datetime, so those values are attached to UTC. A string that
    already carries an offset keeps that offset. The search-log note that
    mentions IST is analyst commentary, not a timezone column, and is not
    applied here.
    """
    try:
        if _is_blank(value):
            if required:
                raise CellProblem(f"row {excel_row}: {field_name} is required")
            return None
        parsed = _coerce_datetime(value, field_name, excel_row)
        if parsed.tzinfo is None:
            return parsed.replace(tzinfo=timezone.utc)
        return parsed
    except CellProblem as exc:
        reasons.append(exc.reason)
        return None


def _coerce_datetime(value: object, field_name: str, excel_row: int) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, bool):
        raise CellProblem(
            f"row {excel_row}: {field_name} is not a valid date ({value!r})"
        )
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
        except (ValueError, OverflowError, TypeError, OSError):
            raise CellProblem(
                f"row {excel_row}: {field_name} is not a valid date ({value!r})"
            ) from None
        if not isinstance(converted, datetime):
            raise CellProblem(
                f"row {excel_row}: {field_name} is not a valid date ({value!r})"
            )
        return converted
    if isinstance(value, str):
        token = value.strip()
        if "/" in token:
            raise CellProblem(
                f"row {excel_row}: {field_name} is not a valid date ({value!r})"
            )
        try:
            return datetime.fromisoformat(token)
        except ValueError:
            raise CellProblem(
                f"row {excel_row}: {field_name} is not a valid date ({value!r})"
            ) from None
    raise CellProblem(
        f"row {excel_row}: {field_name} is not a valid date ({value!r})"
    )


def _take_rating(value: object, excel_row: int, reasons: list[str]) -> float | None:
    try:
        if _is_blank(value):
            return None
        if isinstance(value, bool) or not isinstance(value, (int, float, str)):
            raise CellProblem(f"row {excel_row}: rating must be a number ({value!r})")
        try:
            rating = float(value)
        except ValueError:
            raise CellProblem(
                f"row {excel_row}: rating must be a number ({value!r})"
            ) from None
        if rating < 0:
            raise CellProblem(f"row {excel_row}: rating must be >= 0 ({rating!r})")
        return rating
    except CellProblem as exc:
        reasons.append(exc.reason)
        return None


def _take_whole_number(
    value: object, field_name: str, excel_row: int, reasons: list[str]
) -> int | None:
    try:
        if isinstance(value, bool) or _is_blank(value):
            raise CellProblem(f"row {excel_row}: {field_name} is required")
        if isinstance(value, int):
            number = value
        elif isinstance(value, float) and value.is_integer():
            number = int(value)
        elif isinstance(value, str) and value.strip().isdigit():
            number = int(value.strip())
        else:
            raise CellProblem(
                f"row {excel_row}: {field_name} must be a whole number ({value!r})"
            )
        if number < 0:
            raise CellProblem(f"row {excel_row}: {field_name} must be >= 0 ({number!r})")
        return number
    except CellProblem as exc:
        reasons.append(exc.reason)
        return None


def _remember_author(
    value: object,
    excel_row: int,
    reasons: list[str],
    author_names: set[str],
) -> str | None:
    """Hash input, or null. The name is recorded only so later text can be scrubbed.

    A blank cell is a missing author, not an anonymous placeholder and not a
    guess. A non-text cell is rejected without quoting the cell: the reason
    would otherwise be a copy of the name.
    """
    if isinstance(value, str):
        if value.strip() == "":
            return None
        author_names.add(value)
        if value.strip() != value:
            author_names.add(value.strip())
        return value
    if _is_blank(value):
        return None
    reasons.append(f"row {excel_row}: author_name_raw must be text")
    return None


def _pydantic_reasons(exc: PydanticValidationError, excel_row: int) -> list[str]:
    """Field and message only. The error's ``input`` can be the whole row."""
    reasons: list[str] = []
    for err in exc.errors():
        loc = ".".join(str(part) for part in err["loc"]) or "record"
        reasons.append(f"row {excel_row}: {loc}: {err['msg']}")
    return reasons


def _apply_duplicate_ids(rows: list[_DocumentRow]) -> list[DuplicateSourceItem]:
    grouped: dict[str, list[_DocumentRow]] = {}
    for row in rows:
        if row.source_item_id is not None:
            grouped.setdefault(row.source_item_id, []).append(row)

    duplicates: list[DuplicateSourceItem] = []
    for item_id, group in grouped.items():
        if len(group) < 2:
            continue
        excel_rows = tuple(sorted(row.excel_row for row in group))
        duplicates.append(
            DuplicateSourceItem(source_item_id=item_id, excel_rows=excel_rows)
        )
        listed = ", ".join(str(number) for number in excel_rows)
        for row in group:
            row.reasons.append(
                f"row {row.excel_row}: duplicate source_item_id {item_id!r} "
                f"(rows {listed})"
            )
    duplicates.sort(key=lambda item: item.excel_rows)
    return duplicates


def _apply_parent_links(rows: list[_DocumentRow]) -> None:
    """A parent id must be another row's ``source_item_id``.

    Checked after duplicate rejection, so a child of a rejected parent is
    rejected too. The parent may appear later in the sheet; order is not a
    relationship. A missing parent cell stays null and is not looked up.
    """
    by_id: dict[str, list[_DocumentRow]] = {}
    for row in rows:
        if row.source_item_id is not None:
            by_id.setdefault(row.source_item_id, []).append(row)

    clean_ids = {
        row.source_item_id
        for row in rows
        if row.source_item_id is not None and not row.reasons
    }
    for row in rows:
        parent = row.parent_thread_id
        if parent is None:
            continue
        if row.source_item_id is not None and parent == row.source_item_id:
            row.reasons.append(
                f"row {row.excel_row}: parent_thread_id {parent!r} equals this "
                "row's source_item_id"
            )
            continue
        matches = by_id.get(parent)
        if not matches:
            row.reasons.append(
                f"row {row.excel_row}: parent_thread_id {parent!r} does not "
                "match a source_item_id in this workbook"
            )
            continue
        if parent not in clean_ids:
            parent_rows = ", ".join(str(item.excel_row) for item in matches)
            row.reasons.append(
                f"row {row.excel_row}: parent_thread_id {parent!r} matches "
                f"source_item_id on row {parent_rows}, which was not accepted"
            )


def _repeated_urls(rows: list[_DocumentRow]) -> list[RepeatedSourceUrl]:
    grouped: dict[str, list[_DocumentRow]] = {}
    for row in rows:
        if row.source_url is not None:
            grouped.setdefault(row.source_url, []).append(row)

    repeated: list[RepeatedSourceUrl] = []
    for url, group in grouped.items():
        if len(group) < 2:
            continue
        ordered = sorted(group, key=lambda row: row.excel_row)
        repeated.append(
            RepeatedSourceUrl(
                source_url=url,
                excel_rows=tuple(row.excel_row for row in ordered),
                source_item_ids=tuple(
                    row.source_item_id or "" for row in ordered
                ),
            )
        )
    repeated.sort(key=lambda item: item.excel_rows)
    return repeated


def _warnings(
    documents: list[_DocumentRow],
    search_rows: list[_SearchRow],
    repeated: list[RepeatedSourceUrl],
) -> list[ImportWarning]:
    warnings: list[ImportWarning] = []
    warnings.extend(
        _alias_warnings(documents, DOCUMENT_SHEET, lambda row: row.excel_row, lambda row: row.platform_alias, lambda row: row.source_item_id)
    )
    warnings.extend(
        _alias_warnings(
            search_rows,
            SEARCH_LOG_SHEET,
            lambda row: row.excel_row,
            lambda row: row.platform_alias,
            lambda row: None,
        )
    )
    for item in repeated:
        rows = ", ".join(str(number) for number in item.excel_rows)
        ids = ", ".join(item_id for item_id in item.source_item_ids if item_id)
        warnings.append(
            ImportWarning(
                code="repeated_source_url",
                message=(
                    f"source URL is shared by rows {rows} "
                    f"(source_item_id {ids}). Distinct source_item_id values "
                    "keep the rows, including a thread and a reply that share "
                    "one thread-level URL."
                ),
                sheet=DOCUMENT_SHEET,
                excel_rows=item.excel_rows,
                source_item_ids=item.source_item_ids,
                source_url=item.source_url,
            )
        )

    kept_values = [row.documents_kept for row in search_rows if not row.reasons]
    if kept_values and all(value is not None for value in kept_values):
        kept = sum(value for value in kept_values if value is not None)
        accepted = sum(1 for row in documents if row.document is not None and not row.reasons)
        if kept != accepted:
            warnings.append(
                ImportWarning(
                    code="search_log_kept_count",
                    message=(
                        f"search_log documents_kept sums to {kept}; "
                        f"{accepted} document rows were accepted"
                    ),
                    sheet=SEARCH_LOG_SHEET,
                )
            )
    return warnings


def _alias_warnings(
    rows: list[Any],
    sheet: str,
    row_number: Any,
    alias_of: Any,
    item_of: Any,
) -> list[ImportWarning]:
    aliased = sorted(
        (row for row in rows if alias_of(row)),
        key=row_number,
    )
    if not aliased:
        return []
    by_token: dict[str, list[Any]] = {}
    for row in aliased:
        by_token.setdefault(alias_of(row), []).append(row)
    warnings: list[ImportWarning] = []
    for token, group in by_token.items():
        stored = WORKBOOK_PLATFORM_ALIASES[token]
        excel_rows = tuple(row_number(row) for row in group)
        item_ids = tuple(
            item_id for item_id in (item_of(row) for row in group) if item_id
        )
        warnings.append(
            ImportWarning(
                code="workbook_platform_alias",
                message=(
                    f"source_platform {token!r} is the workbook label for Google "
                    f"Photos Help and is stored as {stored!r}"
                ),
                sheet=sheet,
                excel_rows=excel_rows,
                source_item_ids=item_ids,
            )
        )
    return warnings


def _partition_documents(
    rows: list[_DocumentRow],
) -> tuple[tuple[CollectedDocument, ...], tuple[AcceptedRow, ...], tuple[RowRejection, ...]]:
    documents: list[CollectedDocument] = []
    accepted: list[AcceptedRow] = []
    rejected: list[RowRejection] = []
    for row in sorted(rows, key=lambda item: item.excel_row):
        if row.document is not None and not row.reasons:
            documents.append(row.document)
            accepted.append(
                AcceptedRow(
                    excel_row=row.excel_row,
                    source_item_id=row.document.source_item_id or "",
                    doc_id=row.document.doc_id,
                )
            )
        else:
            rejected.append(
                RowRejection(
                    excel_row=row.excel_row,
                    sheet=DOCUMENT_SHEET,
                    source_item_id=row.source_item_id,
                    reasons=tuple(row.reasons) or (
                        f"row {row.excel_row}: row was not accepted",
                    ),
                )
            )
    return tuple(documents), tuple(accepted), tuple(rejected)


def _empty_report(
    workbook_name: str,
    workbook_sha256: str,
    batch_id: str,
    *,
    structure_errors: tuple[str, ...],
) -> ImportReport:
    return ImportReport(
        workbook_name=workbook_name,
        workbook_sha256=workbook_sha256,
        ingest_batch_id=batch_id,
        rows_read=0,
        rows_accepted=0,
        rows_rejected=0,
        warnings=(),
        duplicate_ids=(),
        repeated_urls=(),
        rejections=(),
        accepted_rows=(),
        search_log_rows_read=0,
        search_log_rows_valid=0,
        search_log_issues=(),
        search_log=(),
        structure_errors=structure_errors,
    )


def _scrub_report(report: ImportReport, names: set[str]) -> ImportReport:
    if not names:
        return report
    payload = report.model_dump()
    return ImportReport.model_validate(_scrub_value(payload, names))


def _scrub_value(value: Any, names: set[str]) -> Any:
    if isinstance(value, str):
        return _scrub_text(value, names)
    if isinstance(value, tuple):
        return tuple(_scrub_value(item, names) for item in value)
    if isinstance(value, list):
        return [_scrub_value(item, names) for item in value]
    if isinstance(value, dict):
        return {key: _scrub_value(item, names) for key, item in value.items()}
    return value


def _scrub_text(text: str, names: set[str]) -> str:
    for name in sorted(names, key=len, reverse=True):
        if name and name in text:
            text = text.replace(name, _REDACTED_AUTHOR)
    return text


def _write_outputs(
    destination: Path,
    report: ImportReport,
    documents: tuple[CollectedDocument, ...],
) -> tuple[Path, Path, Path | None]:
    destination.mkdir(parents=True, exist_ok=True)
    report_json = destination / "import_report.json"
    report_text = destination / "import_report.txt"
    documents_path = destination / "collected_documents.jsonl"

    report_json.write_text(
        json.dumps(report.model_dump(mode="json"), indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    report_text.write_text(_render_report(report), encoding="utf-8")

    if report.structure_errors:
        if documents_path.exists():
            documents_path.unlink()
        return report_json, report_text, None

    lines: list[str] = []
    for document in documents:
        payload = document.model_dump(mode="json")
        _reject_forbidden_keys(payload)
        lines.append(json.dumps(payload, ensure_ascii=False, separators=(",", ":")))
    body = ("\n".join(lines) + "\n") if lines else ""
    documents_path.write_text(body, encoding="utf-8")
    return report_json, report_text, documents_path


def _reject_forbidden_keys(value: object, *, where: str = "collected document") -> None:
    if isinstance(value, dict):
        for key, item in value.items():
            if key in _FORBIDDEN_OUTPUT_KEYS:
                raise ValidationError(
                    f"refusing to write {where}: field {key!r} must not be stored"
                )
            _reject_forbidden_keys(item, where=where)
    elif isinstance(value, list):
        for item in value:
            _reject_forbidden_keys(item, where=where)


def _render_report(report: ImportReport) -> str:
    lines = [
        "Workbook import report",
        f"Workbook: {report.workbook_name}",
        f"SHA-256: {report.workbook_sha256}",
        f"Ingest batch: {report.ingest_batch_id}",
        "",
        "Documents",
        f"  rows read: {report.rows_read}",
        f"  rows accepted: {report.rows_accepted}",
        f"  rows rejected: {report.rows_rejected}",
        "",
        "Warnings",
    ]
    lines.extend(_bullet(report.warnings, _format_warning))
    lines.extend(["", "Duplicate source_item_id"])
    lines.extend(
        _bullet(
            report.duplicate_ids,
            lambda item: (
                f"{item.source_item_id} on rows "
                + ", ".join(str(number) for number in item.excel_rows)
            ),
        )
    )
    lines.extend(["", "Repeated source URLs"])
    lines.extend(
        _bullet(
            report.repeated_urls,
            lambda item: (
                f"rows {', '.join(str(number) for number in item.excel_rows)}: "
                f"{item.source_url}"
            ),
        )
    )
    lines.extend(["", "Row rejections"])
    if report.rejections:
        for rejection in report.rejections:
            item = rejection.source_item_id or "(no source_item_id)"
            lines.append(f"  - row {rejection.excel_row} [{item}]")
            for reason in rejection.reasons:
                lines.append(f"      {reason}")
    else:
        lines.append("  (none)")
    lines.extend(["", "Accepted rows"])
    lines.extend(
        _bullet(
            report.accepted_rows,
            lambda item: f"row {item.excel_row}: {item.source_item_id} -> {item.doc_id}",
        )
    )
    lines.extend(
        [
            "",
            "Search log",
            "  Collection-audit metadata. These rows are not collected documents.",
            f"  rows read: {report.search_log_rows_read}",
            f"  rows valid: {report.search_log_rows_valid}",
        ]
    )
    if report.search_log_issues:
        lines.append("  issues:")
        for issue in report.search_log_issues:
            for reason in issue.reasons:
                lines.append(f"    - {reason}")
    else:
        lines.append("  issues: (none)")
    lines.extend(["", "Structure"])
    lines.extend(_bullet(report.structure_errors, lambda item: item))
    lines.append("")
    return "\n".join(lines)


def _bullet(items: Any, render: Any) -> list[str]:
    if not items:
        return ["  (none)"]
    return [f"  - {render(item)}" for item in items]


def _format_warning(warning: ImportWarning) -> str:
    rows = ""
    if warning.excel_rows:
        rows = " rows " + ", ".join(str(number) for number in warning.excel_rows)
    return f"{warning.sheet}{rows}: {warning.code}: {warning.message}"


def _log_report(report: ImportReport, names: set[str] | None = None) -> None:
    """Log counts and scrubbed reasons. Never the author cell, never the salt."""
    scrub = (lambda text: _scrub_text(text, names)) if names else (lambda text: text)
    _LOG.info(
        "workbook import finished",
        extra={
            "workbook": report.workbook_name,
            "rows_read": report.rows_read,
            "rows_accepted": report.rows_accepted,
            "rows_rejected": report.rows_rejected,
            "warning_count": len(report.warnings),
            "structure_errors": [scrub(item) for item in report.structure_errors],
        },
    )
    for warning in report.warnings:
        _LOG.warning(
            "workbook import warning",
            extra={
                "code": warning.code,
                "sheet": warning.sheet,
                "excel_rows": list(warning.excel_rows),
                "detail": scrub(warning.message),
            },
        )
    for rejection in report.rejections:
        _LOG.warning(
            "document row rejected",
            extra={
                "excel_row": rejection.excel_row,
                "source_item_id": rejection.source_item_id,
                "reasons": [scrub(reason) for reason in rejection.reasons],
            },
        )
    for issue in report.search_log_issues:
        _LOG.warning(
            "search log row rejected",
            extra={
                "excel_row": issue.excel_row,
                "reasons": [scrub(reason) for reason in issue.reasons],
            },
        )
